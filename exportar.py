"""Exportar todos los datos a un Excel (se envía como archivo por Telegram)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import config
import consultas
import datos
from movimientos import CATEGORIAS

HOJAS = [
    ("Saldos", "SELECT fecha AS Fecha, uid AS Cuenta, importe AS Importe FROM saldos ORDER BY fecha, uid"),
    ("Movimientos", "SELECT fecha AS Fecha, uid AS Cuenta, importe AS Importe, contraparte AS Contraparte, "
                    "descripcion AS Concepto, categoria AS Categoria, interno AS Traspaso FROM movimientos ORDER BY fecha"),
    ("Cobros", "SELECT fecha AS Fecha, importe AS Importe, tipo AS Tipo, nota AS Nota FROM ingresos ORDER BY fecha"),
    ("Deudas", "SELECT nombre AS Nombre, CASE sentido WHEN 'me_deben' THEN 'Me deben' ELSE 'Debo' END AS Sentido, "
               "tipo AS Tipo, importe AS Importe, cuota AS Cuota, cuotas_total AS Cuotas, "
               "dia AS Dia, fecha AS Vence, aplazamientos AS Aplazamientos, activa AS Activa FROM deudas WHERE activa >= 0"),
    ("Pagos", "SELECT d.nombre AS Deuda, p.periodo AS Periodo, p.importe AS Importe, p.fecha AS Fecha "
              "FROM pagos p JOIN deudas d ON d.id = p.deuda_id ORDER BY p.fecha"),
    ("Metas", "SELECT nombre AS Nombre, objetivo AS Objetivo, fecha_limite AS Fecha, ahorrado AS Inicial, "
              "uid AS Cuenta, aporte_modo AS Aportacion, aporte_fijo AS Fija, completada AS Conseguida "
              "FROM metas WHERE activa >= 0"),
    ("Aportes por ciclo", "SELECT m.nombre AS Meta, c.ciclo AS Ciclo, c.estado AS Estado, c.importe AS Importe, "
                          "c.fecha AS Fecha FROM aportes_ciclo c JOIN metas m ON m.id = c.meta_id ORDER BY c.ciclo"),
    ("Aportaciones", "SELECT m.nombre AS Meta, a.fecha AS Fecha, a.importe AS Importe FROM aportaciones a "
                     "JOIN metas m ON m.id = a.meta_id ORDER BY a.fecha"),
]


def hoja_proyeccion(libro):
    """Hoja con la proyección cobro a cobro (se calcula al exportar)."""
    import proyeccion
    p = proyeccion.simular()
    if not p:
        return
    hoja = libro.create_sheet("Proyeccion")
    columnas = ["Cobro", "Salario", "Te devuelven", "Gastos fijos", "Suscripciones", "Cuotas", "Pagos unicos",
                "Metas", "Libre", "Acumulado", "Estado"]
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    estados = {"🟢": "Sostenible", "🟡": "Justo", "🔴": "No llega"}
    hoja.append(["Hoy (hasta el cobro)", None, p["deben_hoy"], None, None, None, None, None, p["libre_hoy"], p["puente"], ""])
    for c in p["ciclos"]:
        t = c["totales"]
        hoja.append([c["fecha"].isoformat(), c["salario"], t["deben"], t["fijos"], t["subs"], t["cuotas"], t["unicos"],
                     t["metas"], c["libre"], c["acumulado"], estados[c["estado"]]])
    for n, columna in enumerate(columnas, 1):
        hoja.column_dimensions[get_column_letter(n)].width = max(12, len(columna) + 4)


def excel():
    """(archivo en memoria, nombre)"""
    nombres = {c["uid"]: consultas.nombre_cuenta(c) for c in consultas.cuentas()}
    libro = Workbook()
    libro.remove(libro.active)
    for titulo, sql in HOJAS:
        hoja = libro.create_sheet(titulo)
        columnas, filas = datos.tabla(sql)
        hoja.append(columnas)
        for celda in hoja[1]:
            celda.font = Font(bold=True)
        for fila in filas:
            if "Cuenta" in columnas:
                n = columnas.index("Cuenta")
                fila[n] = nombres.get(fila[n], fila[n])
            if "Categoria" in columnas:
                n = columnas.index("Categoria")
                fila[n] = CATEGORIAS.get(fila[n], fila[n])
            hoja.append(fila)
        for n, columna in enumerate(columnas, 1):
            hoja.column_dimensions[get_column_letter(n)].width = max(12, len(columna) + 4)
    hoja_proyeccion(libro)
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida, f"finanzas_{config.hoy()}.xlsx"
